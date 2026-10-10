import os, re
from pathlib import Path
from datetime import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
os.environ["OPENPYXL_LXML"] = "False"
try:
    import openpyxl
except ImportError:
    import tkinter.messagebox as mb
    mb.showerror("Missing library", "openpyxl is not installed.\n\nRun:\npip install openpyxl")
    raise SystemExit

TITLE = "PDF -> Excel Licence Tool"
PATTERN = re.compile(r"^(?P<yref>[^-]+)-(?P<licence>[^.]+)\.pdf$", re.I)

class App:
    def __init__(self, root):
        self.root = root
        self.root.title(TITLE)
        self.root.geometry("820x600")
        self.wb = None
        self.ws = None
        self.excel = None
        self.pdfs = []
        self.results = []
        self.step()

    def step(self):
        for w in self.root.winfo_children():
            w.destroy()
        outer = ttk.Frame(self.root, padding=24)
        outer.pack(fill="both", expand=True)
        ttk.Label(outer, text=TITLE, font=("TkDefaultFont", 20, "bold")).pack(anchor="w")
        ttk.Label(outer, text="Safely match PDF filenames to an Excel reference column and fill a licence column.", wraplength=740).pack(anchor="w", pady=(5, 20))
        self.body = ttk.Frame(outer)
        self.body.pack(fill="both", expand=True)
        self.buttons = ttk.Frame(outer)
        self.buttons.pack(fill="x", pady=(15, 0))
        self.excel_screen()

    def clear_body(self):
        for w in self.body.winfo_children():
            w.destroy()
        for w in self.buttons.winfo_children():
            w.destroy()

    def excel_screen(self):
        self.clear_body()
        ttk.Label(self.body, text="1. Select Excel workbook", font=("TkDefaultFont", 14, "bold")).pack(anchor="w")
        self.excel_var = tk.StringVar()
        ttk.Entry(self.body, textvariable=self.excel_var, width=90).pack(fill="x", pady=12)
        ttk.Button(self.body, text="Browse for Excel...", command=self.pick_excel).pack(anchor="w")
        ttk.Button(self.buttons, text="Continue", command=self.open_excel).pack(side="right")
        ttk.Button(self.buttons, text="Exit", command=self.root.destroy).pack(side="right", padx=8)

    def pick_excel(self):
        p = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx *.xlsm"), ("All files", "*.*")])
        if p:
            self.excel_var.set(p)

    def open_excel(self):
        p = self.excel_var.get().strip()
        if not os.path.isfile(p):
            return messagebox.showwarning("Excel", "Select a valid Excel file.")
        try:
            self.excel = Path(p)
            self.wb = openpyxl.load_workbook(p, keep_vba=p.lower().endswith(".xlsm"))
        except Exception as e:
            return messagebox.showerror("Excel error", str(e))
        self.sheet_screen()

    def sheet_screen(self):
        self.clear_body()
        ttk.Label(self.body, text="2. Select worksheet", font=("TkDefaultFont", 14, "bold")).pack(anchor="w")
        self.sheet_var = tk.StringVar()
        sheets = self.wb.sheetnames
        preferred = next((s for s in sheets if s.lower() == "customer list"), sheets[0])
        self.sheet_var.set(preferred)
        ttk.Combobox(self.body, textvariable=self.sheet_var, values=sheets, state="readonly", width=70).pack(anchor="w", pady=12)
        ttk.Button(self.buttons, text="Continue", command=self.open_sheet).pack(side="right")
        ttk.Button(self.buttons, text="Back", command=self.excel_screen).pack(side="right", padx=8)
        ttk.Button(self.buttons, text="Exit", command=self.root.destroy).pack(side="right")

    def open_sheet(self):
        self.ws = self.wb[self.sheet_var.get()]
        self.column_screen()

    def column_options(self):
        out = []
        for c in range(1, min(self.ws.max_column, 200) + 1):
            vals = []
            for r in range(1, min(self.ws.max_row, 20) + 1):
                v = self.ws.cell(r, c).value
                if v not in (None, ""):
                    vals.append(str(v).strip())
            header = next(iter(vals), "(blank)")
            out.append((c, f"{openpyxl.utils.get_column_letter(c)} - {header}"))
        return out

    def detect(self, names):
        names = {x.lower() for x in names}
        for c in range(1, min(self.ws.max_column, 200) + 1):
            for r in range(1, min(self.ws.max_row, 20) + 1):
                v = self.ws.cell(r, c).value
                if isinstance(v, str) and v.strip().lower() in names:
                    return c
        return None

    def column_screen(self):
        self.clear_body()
        ttk.Label(self.body, text="3. Map columns", font=("TkDefaultFont", 14, "bold")).pack(anchor="w")
        opts = self.column_options()
        labels = [x[1] for x in opts]
        yc = self.detect(["y ref", "yref", "ref", "reference"])
        lc = self.detect(["licence no.", "licence no", "license no.", "license no", "licence", "license"])
        self.yvar = tk.StringVar(value=next((x[1] for x in opts if x[0] == yc), labels[0]))
        self.lvar = tk.StringVar(value=next((x[1] for x in opts if x[0] == lc), labels[0]))
        f = ttk.Frame(self.body)
        f.pack(fill="x", pady=15)
        ttk.Label(f, text="Y Ref column:").grid(row=0, column=0, sticky="w", pady=8)
        ttk.Combobox(f, textvariable=self.yvar, values=labels, state="readonly", width=65).grid(row=0, column=1, sticky="ew", pady=8)
        ttk.Label(f, text="Licence No. column:").grid(row=1, column=0, sticky="w", pady=8)
        ttk.Combobox(f, textvariable=self.lvar, values=labels, state="readonly", width=65).grid(row=1, column=1, sticky="ew", pady=8)
        f.columnconfigure(1, weight=1)
        ttk.Label(self.body, text="Auto-detect will try to select G - Y ref and J - LICENCE NO.", wraplength=740).pack(anchor="w")
        ttk.Button(self.buttons, text="Continue", command=self.columns_done).pack(side="right")
        ttk.Button(self.buttons, text="Back", command=self.sheet_screen).pack(side="right", padx=8)
        ttk.Button(self.buttons, text="Exit", command=self.root.destroy).pack(side="right")

    def colnum(self, s):
        return openpyxl.utils.column_index_from_string(s.split("-")[0].strip())

    def columns_done(self):
        self.ycol = self.colnum(self.yvar.get())
        self.lcol = self.colnum(self.lvar.get())
        if self.ycol == self.lcol:
            return messagebox.showwarning("Columns", "Choose two different columns.")
        self.folder_screen()

    def folder_screen(self):
        self.clear_body()
        ttk.Label(self.body, text="4. Select PDF folder", font=("TkDefaultFont", 14, "bold")).pack(anchor="w")
        self.folder_var = tk.StringVar()
        ttk.Entry(self.body, textvariable=self.folder_var, width=90).pack(fill="x", pady=12)
        ttk.Button(self.body, text="Browse for folder...", command=self.pick_folder).pack(anchor="w")
        self.count = tk.StringVar()
        ttk.Label(self.body, textvariable=self.count).pack(anchor="w", pady=12)
        ttk.Button(self.buttons, text="Continue", command=self.load_folder).pack(side="right")
        ttk.Button(self.buttons, text="Back", command=self.column_screen).pack(side="right", padx=8)
        ttk.Button(self.buttons, text="Exit", command=self.root.destroy).pack(side="right")

    def pick_folder(self):
        p = filedialog.askdirectory(title="Select PDF folder")
        if p:
            self.folder_var.set(p)
            self.show_count()

    def show_count(self):
        p = self.folder_var.get()
        if os.path.isdir(p):
            n = sum(1 for x in Path(p).iterdir() if x.is_file() and x.suffix.lower() == ".pdf")
            self.count.set(f"PDF files found: {n}")

    def load_folder(self):
        p = self.folder_var.get()
        if not os.path.isdir(p):
            return messagebox.showwarning("Folder", "Select a valid folder.")
        self.pdfs = sorted([x for x in Path(p).iterdir() if x.is_file() and x.suffix.lower() == ".pdf"], key=lambda x: x.name.lower())
        if not self.pdfs:
            return messagebox.showwarning("PDFs", "No PDFs found.")
        self.ready_screen()

    def ready_screen(self):
        self.clear_body()
        ttk.Label(self.body, text="5. Ready to process", font=("TkDefaultFont", 14, "bold")).pack(anchor="w")
        ttk.Label(self.body, text=f"Excel: {self.excel.name}\nSheet: {self.ws.title}\nY Ref: {openpyxl.utils.get_column_letter(self.ycol)}\nLicence: {openpyxl.utils.get_column_letter(self.lcol)}\nPDFs: {len(self.pdfs)}", justify="left").pack(anchor="w", pady=18)
        ttk.Label(self.body, text="Nothing has been changed yet. Existing licence values will never be overwritten automatically.", wraplength=740).pack(anchor="w")
        ttk.Button(self.buttons, text="START PROCESSING", command=self.start).pack(side="right")
        ttk.Button(self.buttons, text="Back", command=self.folder_screen).pack(side="right", padx=8)
        ttk.Button(self.buttons, text="Exit", command=self.root.destroy).pack(side="right")

    def start(self):
        self.results = []
        self.i = 0
        self.lookup = {}
        self.dupes = {}
        for r in range(1, self.ws.max_row + 1):
            v = self.ws.cell(r, self.ycol).value
            if v in (None, ""):
                continue
            k = str(v).strip()
            if k in self.lookup:
                self.dupes.setdefault(k, [self.lookup[k]]).append(r)
            else:
                self.lookup[k] = r
        self.next_pdf()

    def next_pdf(self):
        if self.i >= len(self.pdfs):
            return self.summary()
        pdf = self.pdfs[self.i]
        self.i += 1
        self.clear_body()
        ttk.Label(self.body, text=f"PDF {self.i} of {len(self.pdfs)}", font=("TkDefaultFont", 15, "bold")).pack(anchor="w")
        ttk.Label(self.body, text=pdf.name, font=("TkDefaultFont", 12)).pack(anchor="w", pady=12)
        m = PATTERN.match(pdf.name)
        if not m:
            return self.result(pdf, None, None, "INVALID FILENAME", "Filename must be YREF-LICENCE.pdf.", "No change will be made.")
        yref = m.group("yref").strip()
        lic = m.group("licence").strip()
        ttk.Label(self.body, text=f"Y Ref: {yref}\nLicence No.: {lic}", justify="left").pack(anchor="w", pady=8)
        if yref in self.dupes:
            return self.result(pdf, yref, lic, "DUPLICATE Y REF", f"Found on rows {self.dupes[yref]}.", "The tool will not guess.")
        row = self.lookup.get(yref)
        if row is None:
            return self.result(pdf, yref, lic, "NOT FOUND", "Y ref does not exist in the selected Excel column.", "No change will be made.")
        current = self.ws.cell(row, self.lcol).value
        if current not in (None, ""):
            return self.result(pdf, yref, lic, "ALREADY FILLED", f"Found on row {row}. Existing licence: {current}.", "Existing value will not be overwritten.")
        self.pending = (pdf, yref, lic, row)
        self.result(pdf, yref, lic, "MATCH FOUND", f"Found on Excel row {row}. Licence cell is empty.", f"CONTINUE will enter {lic}.", fill=True)

    def result(self, pdf, yref, lic, status, detail, action, fill=False):
        ttk.Label(self.body, text=status, font=("TkDefaultFont", 18, "bold")).pack(anchor="w", pady=12)
        ttk.Label(self.body, text=detail, wraplength=740).pack(anchor="w", pady=6)
        ttk.Label(self.body, text=action, wraplength=740).pack(anchor="w", pady=6)
        self.current = (pdf, yref, lic, status, detail)
        if fill:
            ttk.Button(self.buttons, text="CONTINUE - FILL", command=self.accept).pack(side="right")
        else:
            ttk.Button(self.buttons, text="CONTINUE", command=self.continue_nonfill).pack(side="right")
        ttk.Button(self.buttons, text="SKIP", command=self.skip).pack(side="right", padx=8)
        ttk.Button(self.buttons, text="EXIT", command=self.root.destroy).pack(side="right")

    def add(self, pdf, yref, lic, status, row="", details=""):
        self.results.append({"PDF": pdf.name, "Y Ref": yref or "", "Licence No.": lic or "", "Result": status, "Excel Row": row, "Details": details})

    def accept(self):
        pdf, yref, lic, row = self.pending
        self.ws.cell(row, self.lcol).value = lic
        self.add(pdf, yref, lic, "FILLED", row, "Licence entered.")
        self.next_pdf()

    def continue_nonfill(self):
        pdf, yref, lic, status, detail = self.current
        row = ""
        if status == "ALREADY FILLED":
            m = re.search(r"row (\d+)", detail)
            row = m.group(1) if m else ""
        self.add(pdf, yref, lic, status, row, detail)
        self.next_pdf()

    def skip(self):
        pdf, yref, lic, status, detail = self.current
        self.add(pdf, yref, lic, "SKIPPED", "", status + ": " + detail)
        self.next_pdf()

    def summary(self):
        self.clear_body()
        ttk.Label(self.body, text="PROCESSING COMPLETE", font=("TkDefaultFont", 18, "bold")).pack(anchor="w", pady=12)
        counts = {}
        for x in self.results:
            counts[x["Result"]] = counts.get(x["Result"], 0) + 1
        ttk.Label(self.body, text=f"PDFs processed: {len(self.results)}\n\n" + "\n".join(f"{k}: {v}" for k, v in sorted(counts.items())), justify="left").pack(anchor="w")
        ttk.Label(self.body, text="\nThe original workbook has not been overwritten. Save Results creates a new updated workbook and a report.", wraplength=740).pack(anchor="w", pady=18)
        ttk.Button(self.buttons, text="SAVE RESULTS", command=self.save).pack(side="right")
        ttk.Button(self.buttons, text="EXIT WITHOUT SAVING", command=self.root.destroy).pack(side="right", padx=8)

    def save(self):
        stem = self.excel.stem
        suffix = self.excel.suffix
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out = self.excel.with_name(stem + "_UPDATED_" + ts + suffix)
        report = self.excel.with_name(stem + "_MATCHING_REPORT_" + ts + ".xlsx")
        try:
            self.wb.save(out)
            rb = openpyxl.Workbook()
            ws = rb.active
            ws.title = "Results"
            headers = ["PDF", "Y Ref", "Licence No.", "Result", "Excel Row", "Details"]
            ws.append(headers)
            for r in self.results:
                ws.append([r[h] for h in headers])
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            s = rb.create_sheet("Summary")
            s.append(["Metric", "Count"])
            s.append(["PDFs processed", len(self.results)])
            counts = {}
            for r in self.results:
                counts[r["Result"]] = counts.get(r["Result"], 0) + 1
            for k, v in sorted(counts.items()):
                s.append([k, v])
            rb.save(report)
        except Exception as e:
            return messagebox.showerror("Save failed", str(e))
        messagebox.showinfo("Saved", f"Updated Excel:\n{out}\n\nReport:\n{report}")
        self.root.destroy()

root = tk.Tk()
App(root)
root.mainloop()